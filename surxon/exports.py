"""Excel export for report specs."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def to_xlsx(spec):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hisobot'
    cols = spec['columns']
    ws.append(['SURXON PAXTA HISOB TIZIMI'])
    ws.append([spec['title']])
    ws.append([])
    ws['A1'].font = Font(bold=True, size=14, color='06335F')
    ws['A2'].font = Font(bold=True, size=12)
    ws.append([c[1] for c in cols])
    head = ws.max_row
    fill = PatternFill('solid', fgColor='06335F')
    thin = Side(style='thin', color='D6E2EE')
    for i in range(1, len(cols) + 1):
        cell = ws.cell(row=head, column=i)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in spec['rows']:
        ws.append([r.get(c[0]) for c in cols])
    for row in ws.iter_rows(min_row=head + 1, max_row=ws.max_row):
        for cell, col in zip(row, cols):
            cell.border = Border(bottom=thin)
            if col[2] == 'kg':
                cell.number_format = '#,##0.0'
            elif col[2] == 'money':
                cell.number_format = '#,##0'
    totals = spec.get('totals') or {}
    if totals:
        ws.append([('Jami' if i == 0 else totals.get(c[0], '')) for i, c in enumerate(cols)])
        for cell, col in zip(ws[ws.max_row], cols):
            cell.font = Font(bold=True)
            if col[2] in ('kg', 'money'):
                cell.number_format = '#,##0.0' if col[2] == 'kg' else '#,##0'
    for i, col in enumerate(cols, start=1):
        width = max([len(str(col[1]))] + [len(str(r.get(col[0]) or '')) for r in spec['rows'][:500]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 45)
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
